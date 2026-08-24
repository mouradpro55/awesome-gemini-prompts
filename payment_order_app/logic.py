import pandas as pd
import openpyxl
import os
import json
import logging
from shutil import copyfile

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

def process_orders(template_path, data_path, output_dir, config_path, progress_callback=None, log_callback=None):
    """
    Process the excel data and generate populated templates.
    progress_callback: func(current, total)
    log_callback: func(message)
    """
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    try:
        with open(config_path, 'r', encoding='utf-8') as f:
            config = json.load(f)
    except Exception as e:
        msg = f"خطأ في قراءة ملف الإعدادات: {str(e)}"
        if log_callback: log_callback(msg)
        return False, msg

    try:
        df = pd.read_excel(data_path)
    except Exception as e:
        msg = f"خطأ في قراءة ملف البيانات: {str(e)}"
        if log_callback: log_callback(msg)
        return False, msg

    sheet_name = config.get("sheet_name", "حوالة")
    mapping = config.get("mapping", {})
    filename_columns = config.get("filename_columns", [])

    total_rows = len(df)

    if total_rows == 0:
        msg = "ملف البيانات فارغ."
        if log_callback: log_callback(msg)
        return False, msg

    for index, row in df.iterrows():
        try:
            # Check for missing values in mandatory columns (all mapped columns for now)
            missing_cols = []
            for col_name in mapping.keys():
                if col_name in row and pd.isna(row[col_name]):
                    missing_cols.append(col_name)

            if missing_cols:
                msg = f"تنبيه: خلايا إجبارية فارغة في الصف {index + 1}: {', '.join(missing_cols)}"
                if log_callback: log_callback(msg)

            # Prepare filename
            filename_parts = ["حوالة"]
            for col in filename_columns:
                if col in row:
                    val = str(row[col]).strip()
                    # Safe filename string
                    val = "".join([c for c in val if c.isalnum() or c in (' ', '_', '-')])
                    if val:
                        filename_parts.append(val)

            out_filename = "_".join(filename_parts) + ".xlsx"
            out_filepath = os.path.join(output_dir, out_filename)

            # Load template using openpyxl directly
            wb = openpyxl.load_workbook(template_path)

            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb.active # fallback to active if name not found
                if log_callback: log_callback(f"تنبيه: ورقة العمل '{sheet_name}' غير موجودة، تم استخدام الورقة الحالية.")

            # Map values
            for col_name, cell_coord in mapping.items():
                if col_name in row:
                    val = row[col_name]
                    # Handle NaNs
                    if pd.isna(val):
                        val = ""
                    ws[cell_coord].value = val

            wb.save(out_filepath)

            msg = f"تم بنجاح توليد: {out_filename}"
            if log_callback: log_callback(msg)

        except Exception as e:
            msg = f"خطأ في معالجة الصف {index + 1}: {str(e)}"
            if log_callback: log_callback(msg)

        if progress_callback:
            progress_callback(index + 1, total_rows)

    return True, "تم الانتهاء من توليد جميع الحوالات."
