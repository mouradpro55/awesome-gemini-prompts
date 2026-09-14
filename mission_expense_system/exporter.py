import openpyxl
from openpyxl.styles import Alignment, Font, Border, Side
import os
from tafqeet import tafqeet

def create_template_if_not_exists(template_path="Template.xlsx"):
    if os.path.exists(template_path):
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "كشف مصاريف التنقل"
    ws.sheet_view.rightToLeft = True

    # Styling helpers
    bold_font = Font(name="Arial", size=12, bold=True)
    normal_font = Font(name="Arial", size=11)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))

    def set_cell(coord, value, font, merge=None):
        ws[coord] = value
        ws[coord].font = font
        ws[coord].alignment = center_align
        if merge:
            ws.merge_cells(merge)

    # Header
    set_cell("A1", "الجمهورية الجزائرية الديمقراطية الشعبية", bold_font, "A1:J1")
    set_cell("A2", "وزارة البريد والمواصلات السلكية واللاسلكية", bold_font, "A2:J2")
    set_cell("A3", "كشف مصاريف التنقل", Font(name="Arial", size=16, bold=True), "A3:J3")

    # Employee Info
    set_cell("A5", "الاسم واللقب:", bold_font)
    ws['B5'].font = normal_font
    set_cell("D5", "الرتبة / الوظيفة:", bold_font)
    ws['E5'].font = normal_font
    set_cell("G5", "الرقم الاستدلالي:", bold_font)
    ws['H5'].font = normal_font

    set_cell("A6", "المقر الإداري:", bold_font)
    ws['B6'].font = normal_font
    set_cell("D6", "طبيعة الحساب:", bold_font)
    ws['E6'].font = normal_font
    set_cell("G6", "رقم الحساب:", bold_font)
    ws['H6'].font = normal_font

    # Missions Table Header
    headers = ["رقم وتاريخ الأمر", "المنطقة", "وسيلة النقل", "تاريخ الذهاب", "تاريخ الإياب",
               "عدد الوجبات", "المبلغ للوجبات", "عدد الليالي", "المبلغ لليالي", "المجموع الفرعي"]

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=8, column=col_idx)
        cell.value = header
        cell.font = bold_font
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 15

    # Footer placeholders
    set_cell("A18", "المجموع العام:", bold_font, "A18:B18")
    set_cell("C18", "", normal_font, "C18:D18") # for the total numeric
    set_cell("E18", "المبلغ بالحروف:", bold_font, "E18:F18")
    set_cell("G18", "", normal_font, "G18:J18") # for tafqeet

    set_cell("B20", "إمضاء المعني بالأمر", bold_font, "B20:C20")
    set_cell("H20", "تأشيرة السيد المدير الآمر بالصرف", bold_font, "H20:J20")

    # Author
    set_cell("A25", "تم إعداد هذا البرنامج من طرف السيد مصباح مراد", Font(name="Arial", size=10, italic=True), "A25:J25")

    wb.save(template_path)

def generate_report(employee, missions, output_path):
    template_path = "Template.xlsx"
    create_template_if_not_exists(template_path)

    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    # Fill Employee Info
    ws['B5'] = employee[1] # full_name
    ws['E5'] = f"{employee[2]} / {employee[3]}" # rank / job_title
    ws['H5'] = employee[4] # index_number
    ws['B6'] = employee[6] # admin_location
    ws['E6'] = employee[7] # bank_type
    ws['H6'] = employee[8] # bank_account

    # Fill Missions
    start_row = 9
    total_meals_amount = 0
    total_nights_amount = 0
    total_general = 0

    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    normal_font = Font(name="Arial", size=11)

    for i, m in enumerate(missions):
        row = start_row + i

        rate = 800 if m[9] == "شمال" else 1000
        n_rate = 3200 if m[9] == "شمال" else 4000

        meals_amt = m[10] * rate
        nights_amt = m[11] * n_rate
        sub_tot = meals_amt + nights_amt

        total_meals_amount += meals_amt
        total_nights_amount += nights_amt
        total_general += sub_tot

        data_row = [
            f"{m[2]}\n{m[3]}", # order num / date
            m[9], # region
            m[6], # transport
            m[7], # dep
            m[8], # ret
            m[10], # meals_c
            meals_amt,
            m[11], # nights_c
            nights_amt,
            sub_tot
        ]

        for col_idx, val in enumerate(data_row, 1):
            cell = ws.cell(row=row, column=col_idx)
            cell.value = val
            cell.font = normal_font
            cell.alignment = center_align
            cell.border = thin_border

    # Adjust Footer position based on number of missions
    footer_start_row = max(18, start_row + len(missions) + 2)

    ws.cell(row=footer_start_row, column=1).value = "المجموع العام:"
    ws.cell(row=footer_start_row, column=1).font = Font(name="Arial", size=12, bold=True)
    ws.cell(row=footer_start_row, column=3).value = f"{total_general} دج"
    ws.cell(row=footer_start_row, column=3).font = normal_font

    ws.cell(row=footer_start_row, column=5).value = "المبلغ بالحروف:"
    ws.cell(row=footer_start_row, column=5).font = Font(name="Arial", size=12, bold=True)
    ws.cell(row=footer_start_row, column=7).value = tafqeet(total_general)
    ws.cell(row=footer_start_row, column=7).font = normal_font
    ws.merge_cells(start_row=footer_start_row, start_column=7, end_row=footer_start_row, end_column=10)

    # Signatures
    ws.cell(row=footer_start_row+2, column=2).value = "إمضاء المعني بالأمر"
    ws.cell(row=footer_start_row+2, column=2).font = Font(name="Arial", size=12, bold=True)

    ws.cell(row=footer_start_row+2, column=8).value = "تأشيرة السيد المدير الآمر بالصرف"
    ws.cell(row=footer_start_row+2, column=8).font = Font(name="Arial", size=12, bold=True)

    # Author string
    ws.cell(row=footer_start_row+6, column=1).value = "تم إعداد هذا البرنامج من طرف السيد مصباح مراد"
    ws.cell(row=footer_start_row+6, column=1).font = Font(name="Arial", size=10, italic=True)
    ws.merge_cells(start_row=footer_start_row+6, start_column=1, end_row=footer_start_row+6, end_column=10)
    ws.cell(row=footer_start_row+6, column=1).alignment = Alignment(horizontal="center")

    wb.save(output_path)
    return output_path

if __name__ == "__main__":
    create_template_if_not_exists()
    print("Template created/verified.")
