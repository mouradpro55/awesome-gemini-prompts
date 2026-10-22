import openpyxl
import os
from allowance_rules import ALLOWANCE_TABLE

def generate_excel_report(employee, missions, output_path, tafqeet_text):
    template_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Template.xlsx")

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"ملف القالب الأصلي غير موجود في المسار: {template_path}. يرجى وضع Template.xlsx بجانب البرنامج.")

    wb = openpyxl.load_workbook(template_path)
    ws = wb.active

    def safe_write(coord_str, val):
        coord = openpyxl.utils.coordinate_to_tuple(coord_str)
        for merged_range in ws.merged_cells.ranges:
            if coord[0] >= merged_range.min_row and coord[0] <= merged_range.max_row and \
               coord[1] >= merged_range.min_col and coord[1] <= merged_range.max_col:
                top_left = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
                top_left.value = val
                return top_left
        cell = ws[coord_str]
        cell.value = val
        return cell


    # 1. Fill Employee Info
    safe_write("C11", employee['full_name'])
    safe_write("G11", missions[0]['report_month_year'] if missions and len(missions) > 0 else "")
    safe_write("C12", employee['rank'])
    safe_write("G12", employee['job_title']) # Replaced formula with job_title
    safe_write("C14", employee['index_number'])
    safe_write("G14", employee['category'])
    safe_write("C15", employee['workplace'])
    safe_write("G15", employee['department']) # Added department to G15 based on standard templates
    safe_write("G16", employee['bank_account'])

    total_n_meals = 0
    total_n_nights = 0
    total_s_meals = 0
    total_s_nights = 0

    # 2. Fill Missions
    start_row = 9
    for idx, m in enumerate(missions[:10]):
        row_dep = start_row + (idx * 2)
        row_ret = row_dep + 1

        n_meals, n_nights, s_meals, s_nights = 0, 0, 0, 0
        if m['region'] == "شمال":
            n_meals = m['meals_count']
            n_nights = m['nights_count']
            total_n_meals += n_meals
            total_n_nights += n_nights
        else:
            s_meals = m['meals_count']
            s_nights = m['nights_count']
            total_s_meals += s_meals
            total_s_nights += s_nights

        try:
            dep_parts = str(m['departure_datetime']).split(' ')
            dep_date = dep_parts[0]
            dep_time = dep_parts[1] if len(dep_parts)>1 else ""

            ret_parts = str(m['return_datetime']).split(' ')
            ret_date = ret_parts[0]
            ret_time = ret_parts[1] if len(ret_parts)>1 else ""
        except:
            dep_date, dep_time, ret_date, ret_time = "", "", "", ""

        # Departure Row (row_dep)
        safe_write(f"I{row_dep}", m['purpose']) # Purpose
        safe_write(f"J{row_dep}", m['itinerary']) # Itinerary
        safe_write(f"K{row_dep}", dep_date)
        safe_write(f"L{row_dep}", dep_time)
        safe_write(f"M{row_dep}", m['transport_mode']) # Transport
        safe_write(f"N{row_dep}", n_meals if n_meals > 0 else "")
        safe_write(f"O{row_dep}", n_nights if n_nights > 0 else "")
        safe_write(f"P{row_dep}", s_meals if s_meals > 0 else "")
        safe_write(f"Q{row_dep}", s_nights if s_nights > 0 else "")
        safe_write(f"R{row_dep}", m['order_number']) # Order Num
        safe_write(f"S{row_dep}", f'=IF(K{row_dep}="","","/")')
        safe_write(f"T{row_dep}", f'=IF(K{row_dep}="","",YEAR(DATEVALUE(K{row_dep})))')

        # Return Row (row_ret)
        safe_write(f"K{row_ret}", ret_date)
        safe_write(f"L{row_ret}", ret_time)
        safe_write(f"R{row_ret}", m['order_date']) # Order Date

    # 3. Write dynamic calculations based on employee category
    category = employee['category']
    cat_rates = ALLOWANCE_TABLE.get(category, ALLOWANCE_TABLE["1-10"])

    rate_n_m = cat_rates.get("شمال", {}).get("meal", 800)
    rate_n_n = cat_rates.get("شمال", {}).get("night", 3200)
    rate_s_m = cat_rates.get("جنوب", {}).get("meal", 1000)
    rate_s_n = cat_rates.get("جنوب", {}).get("night", 4000)

    safe_write("N29", total_n_meals if total_n_meals > 0 else "")
    safe_write("O29", total_n_nights if total_n_nights > 0 else "")
    safe_write("P29", total_s_meals if total_s_meals > 0 else "")
    safe_write("Q29", total_s_nights if total_s_nights > 0 else "")

    safe_write("D19", total_n_meals if total_n_meals > 0 else "")
    safe_write("D20", total_n_nights if total_n_nights > 0 else "")
    safe_write("D22", total_s_meals if total_s_meals > 0 else "")
    safe_write("D23", total_s_nights if total_s_nights > 0 else "")

    safe_write("G19", (total_n_meals * rate_n_m) if total_n_meals > 0 else "")
    safe_write("G20", (total_n_nights * rate_n_n) if total_n_nights > 0 else "")
    safe_write("G22", (total_s_meals * rate_s_m) if total_s_meals > 0 else "")
    safe_write("G23", (total_s_nights * rate_s_n) if total_s_nights > 0 else "")

    grand_total = (total_n_meals * rate_n_m) + (total_n_nights * rate_n_n) + (total_s_meals * rate_s_m) + (total_s_nights * rate_s_n)
    safe_write("G25", grand_total if grand_total > 0 else "")

    # 4. Tafqeet Text
    if grand_total > 0:
        if not tafqeet_text.endswith("دينار جزائري"):
            tafqeet_text += " دينار جزائري"
        safe_write("B29", f"أوقِف هذا الكشف عند مبلغ قدره: {tafqeet_text} تماماً")
    else:
        safe_write("B29", "")

    wb.save(output_path)
