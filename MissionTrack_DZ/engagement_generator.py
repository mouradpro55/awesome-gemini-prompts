import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_rtl(p):
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement('w:bidi'))

def set_cell_background(cell, fill_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def set_table_rtl(table):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        bidi = OxmlElement('w:bidiVisual')
        tblPr[0].append(bidi)

def set_cell_borders(cell):
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        b = OxmlElement(f'w:{border_name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '000000')
        tcBorders.append(b)
    tcPr.append(tcBorders)

def set_table_borders(table):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = OxmlElement('w:tblBorders')
        for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '4')
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), '000000')
            borders.append(b)
        tblPr[0].append(borders)

def generate_engagement_doc(output_path, employee, proposed_amount, prior_commitments, budget_info, commitment_num="01", commitment_date="06-04-2026", allocated_ae=None):
    """
    توليد بطاقة الالتزام المالي مطابقة للنموذج الرسمي
    الحساب المالي:
      الرصيد الأولي = رخصة الالتزام - الالتزامات السابقة
      الرصيد المتبقي = الرصيد الأولي - الالتزام المقترح
    """
    allocated_ae = float(allocated_ae) if allocated_ae is not None else (budget_info[14] if len(budget_info) > 14 else 600000.00)
    initial_balance = allocated_ae - prior_commitments
    remaining_balance = initial_balance - proposed_amount

    doc = docx.Document()

    # ضبط الهوامش الأصلية للنموذج (Top: 1.5 cm, Others: 2.5 cm)
    s = doc.sections[0]
    s.top_margin = Inches(0.59)
    s.bottom_margin = Inches(0.98)
    s.left_margin = Inches(0.98)
    s.right_margin = Inches(0.98)

    # 1. الترويسة الإدارية
    p0 = doc.add_paragraph()
    set_rtl(p0)
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run("الجمهورية الجزائرية الديمقراطية الشعبية")
    r0.font.name = 'Arial'; r0.font.size = Pt(12); r0.bold = True

    lines = [
        ("الوزارة : ", "وزارة الداخلية و الجماعات المحلية"),
        ("الهيئة الإدارية : ", "مديرية المواصلات السلكية واللاسلكية الوطنية - المينعة"),
        ("رمز الامر بالصرف : ", budget_info[8] if len(budget_info)>8 else "458/107"),
    ]
    for lbl, val in lines:
        p = doc.add_paragraph()
        set_rtl(p)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(lbl); r1.bold = True; r1.font.name = 'Arial'; r1.font.size = Pt(10.5)
        r2 = p.add_run(val); r2.font.name = 'Arial'; r2.font.size = Pt(10.5)

    # رقم البطاقة والتاريخ
    p_card = doc.add_paragraph()
    set_rtl(p_card)
    p_card.paragraph_format.space_before = Pt(4)
    p_card.paragraph_format.space_after = Pt(4)
    r = p_card.add_run(f"        رقم بطاقة الالتزام : {commitment_num}                                          التاريخ : {commitment_date}")
    r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(10.5)

    codes = [
        ("رمز البرنامج: ", f"{budget_info[2]}                                       {budget_info[3]}"),
        ("رمز النشــاط: ", f"{budget_info[6]}                                    {budget_info[7]}"),
        ("رمز النشاط الفرعي : ", ""),
        ("رمز البرنامج الفرعي: ", f"{budget_info[4]}                      {budget_info[5]}"),
        ("العنوان : ", "")
    ]
    for lbl, val in codes:
        p = doc.add_paragraph()
        set_rtl(p)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(lbl); r1.bold = True; r1.font.name = 'Arial'; r1.font.size = Pt(10)
        r2 = p.add_run(val); r2.bold = True; r2.font.name = 'Arial'; r2.font.size = Pt(10)

    # 2. جداول العناوين (مع تأشير العنوان الثاني بـ X في مربع منفصل)
    for title_text, mark in [("العنوان الأول : نفقات المستخدمين", ""),
                             ("العنوان الثاني : نفقات تسيير المصالح", "X"),
                             ("العنوان الرابع : نفقات التحويل", "")]:
        t = doc.add_table(rows=1, cols=2)
        set_table_rtl(t)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER

        # الخلية 0: مربع الاختيار Checkbox (صغير ومحاط بحدود، يظهر يميناً في أول السطر بفضل RTL)
        cell_m = t.rows[0].cells[0]
        cell_m.width = Inches(0.3)
        set_cell_borders(cell_m)
        p_m = cell_m.paragraphs[0]; set_rtl(p_m); p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_m = p_m.add_run(mark); r_m.bold = True; r_m.font.name = 'Arial'; r_m.font.size = Pt(11)

        # الخلية 1: النص (عريض وبدون حدود، يظهر يسار المربع)
        cell_t = t.rows[0].cells[1]
        cell_t.width = Inches(5.5)
        p_t = cell_t.paragraphs[0]; set_rtl(p_t)
        r_t = p_t.add_run(title_text); r_t.bold = True; r_t.font.name = 'Arial'; r_t.font.size = Pt(10)

        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 3. جدول الالتزام والحساب المالي (Table 3)
    t_fin = doc.add_table(rows=3, cols=7)
    set_table_rtl(t_fin)
    t_fin.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_fin)

    # دمج ترويسة الجدول وتعيين النصوص
    cell_0 = t_fin.cell(0, 0)
    cell_1 = t_fin.cell(0, 1)
    cell_0.merge(cell_1)

    headers = [
        "الصنف / الصنف الفرعي", "رخصة الإلتزام\nالمفتوحة / المعدلة",
        "مجموع الإلتزامات\nالسابقة", "الرصيد الأولي", "الالتزام المقترح", "الرصيد المتبقي"
    ]

    col_indices = [0, 2, 3, 4, 5, 6]
    for i, h in zip(col_indices, headers):
        cell = t_fin.cell(0, i)
        set_cell_background(cell, 'F1F5F9')
        p = cell.paragraphs[0]; set_rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h); r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(8.5)

    # صف نفقات التنقلات (21000)
    account_code = budget_info[10] if budget_info and len(budget_info) > 10 and budget_info[10] else "21000"
    account_label = budget_info[11] if budget_info and len(budget_info) > 11 and budget_info[11] else "التنقلات و النقل و الاتصالات"

    row1_vals = [
        account_code,
        account_label,
        f"{allocated_ae:,.2f}".replace(',', ' '),
        f"{prior_commitments:,.2f}".replace(',', ' '),
        f"{initial_balance:,.2f}".replace(',', ' '),
        f"{proposed_amount:,.2f}".replace(',', ' '),
        f"{remaining_balance:,.2f}".replace(',', ' ')
    ]
    for i, val in enumerate(row1_vals):
        cell = t_fin.cell(1, i)
        p = cell.paragraphs[0]; set_rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(val)); r.font.name = 'Arial'; r.font.size = Pt(9)
        if i >= 2: r.bold = True

    # صف مصاريف الكهرباء والغاز والماء (27800)
    row2_vals = ["27800", "مصاريف الكهرباء و الغاز و الماء", "/", "/", "/", "/", "/"]
    for i, val in enumerate(row2_vals):
        cell = t_fin.cell(2, i)
        p = cell.paragraphs[0]; set_rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val); r.font.name = 'Arial'; r.font.size = Pt(9)

    # 4. موضوع الالتزام
    p_topic = doc.add_paragraph()
    set_rtl(p_topic)
    p_topic.paragraph_format.space_before = Pt(8)
    p_topic.paragraph_format.space_after = Pt(2)
    r_top = p_topic.add_run("موضوع الالتزام :")
    r_top.bold = True; r_top.font.name = 'Arial'; r_top.font.size = Pt(10.5)

    p_body = doc.add_paragraph()
    set_rtl(p_body)
    r_body = p_body.add_run(f"التزام بسند رقم {commitment_num} بتاريخ {commitment_date} - كشف مصاريف تنقل ومهمات لفائدة السيد(ة): {employee[1]} ({employee[2]})")
    r_body.bold = True; r_body.font.name = 'Arial'; r_body.font.size = Pt(10)

    # 5. جدول التأشيرة والإمضاءات (Table 4)
    t_sign = doc.add_table(rows=2, cols=2)
    set_table_rtl(t_sign)
    t_sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_sign)
    t_sign.rows[0].cells[0].width = Inches(3.2)
    t_sign.rows[0].cells[1].width = Inches(3.2)

    p_s0 = t_sign.cell(0, 0).paragraphs[0]; set_rtl(p_s0); p_s0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_s0.add_run("إطار مخصص للمراقب الميزانياتي"); r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(10)

    p_s1 = t_sign.cell(0, 1).paragraphs[0]; set_rtl(p_s1); p_s1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_s1.add_run("إطار مخصص للآمر بالصرف"); r.bold = True; r.font.name = 'Arial'; r.font.size = Pt(10)

    p_box0 = t_sign.cell(1, 0).paragraphs[0]; set_rtl(p_box0)
    p_box0.add_run("رقم التأشيرة :\nتاريخ التأشيرة :\n\nإمضاء :                                  الختم :").font.name = 'Arial'

    p_box1 = t_sign.cell(1, 1).paragraphs[0]; set_rtl(p_box1)
    p_box1.add_run("\nختم \n\nامضاء\n\n").font.name = 'Arial'

    doc.save(output_path)
